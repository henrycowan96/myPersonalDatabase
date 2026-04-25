import os
import sys
from pathlib import Path
from datetime import datetime

# Add parent directory to path for imports
sys.path.append(str(Path(__file__).parent.parent))

# Configuration
DRIVE_FILES_DIR = Path(__file__).parent.parent / "data" / "raw_docs" / "drive_files"

def convert_docx_to_text(docx_path):
    """Convert DOCX file to plain text"""
    try:
        from docx import Document
        
        doc = Document(docx_path)
        text_content = []
        
        # Extract paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)
        
        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_content.append(" | ".join(row_text))
        
        result = "\n".join(text_content)
        if not result or len(result) < 10:
            print(f"Warning: {docx_path.name} appears to be empty or very short")
            return result
        return result
    except ImportError:
        print("Error: python-docx not installed. Run: pip install python-docx")
        return None
    except Exception as e:
        print(f"Error converting {docx_path.name}: {e}")
        return None

def convert_pptx_to_text(pptx_path):
    """Convert PPTX file to plain text"""
    try:
        from pptx import Presentation
        
        prs = Presentation(pptx_path)
        text_content = []
        
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = [f"\n--- Slide {slide_num} ---"]
            
            # Extract text from shapes
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            
            # Extract text from tables
            for shape in slide.shapes:
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = []
                        for cell in row.cells:
                            if cell.text.strip():
                                row_text.append(cell.text.strip())
                        if row_text:
                            slide_text.append(" | ".join(row_text))
            
            if len(slide_text) > 1:  # More than just the slide header
                text_content.extend(slide_text)
        
        return "\n".join(text_content)
    except ImportError:
        print("Error: python-pptx not installed. Run: pip install python-pptx")
        return None
    except Exception as e:
        print(f"Error converting {pptx_path.name}: {e}")
        return None

def convert_xlsx_to_text(xlsx_path):
    """Convert XLSX file to plain text"""
    try:
        import openpyxl
        
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        text_content = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_content.append(f"\n=== Sheet: {sheet_name} ===")
            
            for row in sheet.iter_rows(values_only=True):
                # Filter out None values and convert to string
                row_values = [str(cell) if cell is not None else "" for cell in row]
                # Only add row if it has content
                if any(val.strip() for val in row_values):
                    text_content.append(" | ".join(row_values))
        
        return "\n".join(text_content)
    except ImportError:
        print("Error: openpyxl not installed. Run: pip install openpyxl")
        return None
    except Exception as e:
        print(f"Error converting {xlsx_path.name}: {e}")
        return None

def convert_xls_to_text(xls_path):
    """Convert XLS file to plain text using xlrd or handle XML spreadsheet format"""
    # First check if it's actually an XML spreadsheet (SpreadsheetML)
    try:
        with open(xls_path, 'rb') as f:
            header = f.read(100)
            if b'<?xml' in header or b'<ss:Workbook' in header:
                # It's an XML spreadsheet, use xml parsing
                return convert_xml_spreadsheet(xls_path)
    except:
        pass
    
    # Try xlrd for binary XLS
    try:
        import xlrd
        
        wb = xlrd.open_workbook(xls_path)
        text_content = []
        
        for sheet_num in range(wb.nsheets):
            sheet = wb.sheet_by_index(sheet_num)
            text_content.append(f"\n=== Sheet: {sheet.name} ===")
            
            for row_num in range(sheet.nrows):
                row = sheet.row(row_num)
                row_values = [str(cell.value) if cell.value else "" for cell in row]
                if any(val.strip() for val in row_values):
                    text_content.append(" | ".join(row_values))
        
        return "\n".join(text_content)
    except ImportError:
        print("Error: xlrd not installed. Run: pip install xlrd")
        return None
    except Exception as e:
        print(f"Error converting {xls_path.name}: {e}")
        return None

def convert_xml_spreadsheet(xml_path):
    """Convert XML SpreadsheetML format to plain text"""
    try:
        import xml.etree.ElementTree as ET
        
        # Read file and handle BOM
        with open(xml_path, 'r', encoding='utf-8-sig') as f:
            content = f.read()
        
        tree = ET.ElementTree(ET.fromstring(content))
        root = tree.getroot()
        
        # Handle namespace
        ns = {'ss': 'urn:schemas-microsoft-com:office:spreadsheet'}
        
        text_content = []
        
        # Find all worksheets
        worksheets = root.findall('.//ss:Worksheet', ns)
        
        for worksheet in worksheets:
            sheet_name = worksheet.get('{urn:schemas-microsoft-com:office:spreadsheet}Name', 'Sheet')
            text_content.append(f"\n=== Sheet: {sheet_name} ===")
            
            # Find all rows
            rows = worksheet.findall('.//ss:Row', ns)
            for row in rows:
                row_values = []
                cells = row.findall('.//ss:Cell', ns)
                for cell in cells:
                    # Find data element
                    data = cell.find('.//ss:Data', ns)
                    if data is not None and data.text:
                        row_values.append(data.text.strip())
                    else:
                        row_values.append("")
                
                if any(val.strip() for val in row_values):
                    text_content.append(" | ".join(row_values))
        
        return "\n".join(text_content)
    except Exception as e:
        print(f"Error converting XML spreadsheet {xml_path.name}: {e}")
        return None

def convert_file(file_path):
    """Convert a file to text based on its extension"""
    suffix = file_path.suffix.lower()
    
    # Skip already converted files
    if file_path.stem.endswith('_converted'):
        return None
    
    # Determine output path
    output_path = file_path.parent / f"{file_path.stem}_converted.txt"
    
    # Skip if already converted
    if output_path.exists():
        print(f"Skipping {file_path.name} (already converted)")
        return None
    
    print(f"Converting {file_path.name}...")
    
    if suffix == '.docx':
        text = convert_docx_to_text(file_path)
    elif suffix == '.pptx':
        text = convert_pptx_to_text(file_path)
    elif suffix == '.xlsx':
        text = convert_xlsx_to_text(file_path)
    elif suffix == '.xls':
        text = convert_xls_to_text(file_path)
    else:
        return None
    
    if text:
        # Add metadata header
        header = f"Source: {file_path.name}\n"
        header += f"Converted: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        header += "=" * 50 + "\n\n"
        
        full_text = header + text
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(full_text)
        
        print(f"  -> Saved to {output_path.name}")
        return output_path
    
    return None

def main():
    print("Converting binary files in drive_files directory...")
    print("=" * 50)
    
    if not DRIVE_FILES_DIR.exists():
        print(f"Directory not found: {DRIVE_FILES_DIR}")
        return
    
    # Find all convertible files
    convertible_extensions = {'.docx', '.pptx', '.xlsx', '.xls'}
    files_to_convert = []
    
    for file_path in DRIVE_FILES_DIR.iterdir():
        if file_path.is_file() and file_path.suffix.lower() in convertible_extensions:
            # Skip already converted files
            if not file_path.stem.endswith('_converted'):
                files_to_convert.append(file_path)
    
    if not files_to_convert:
        print("No convertible files found.")
        return
    
    print(f"Found {len(files_to_convert)} files to convert:\n")
    for f in files_to_convert:
        print(f"  - {f.name} ({f.suffix})")
    
    confirm = input(f"\nConvert all {len(files_to_convert)} files? (y/n): ")
    if confirm.lower() != 'y':
        print("Aborted.")
        return
    
    print("\nConverting files...")
    success_count = 0
    error_count = 0
    
    for file_path in files_to_convert:
        try:
            result = convert_file(file_path)
            if result:
                success_count += 1
            else:
                error_count += 1
        except Exception as e:
            print(f"Error processing {file_path.name}: {e}")
            error_count += 1
    
    print(f"\nComplete! Converted {success_count} files to text format.")
    if error_count > 0:
        print(f"Failed to convert {error_count} files")
    print(f"Converted files are saved with '_converted.txt' suffix in {DRIVE_FILES_DIR}")
    print("You can now run 'python scripts/ingest.py' to process these files.")

if __name__ == "__main__":
    main()
